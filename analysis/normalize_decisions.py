"""规范化复核结果：空→valid, i→invalid, u→unsure"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import math
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

SRC = "outputs/screening_v3/tisp_v3_review_queue.xlsx"
CSV = "outputs/screening_v3/tisp_v3_review_queue.csv"

df = pd.read_excel(SRC, sheet_name="复核队列")

def normalize(x):
    if isinstance(x, str):
        x = x.strip()
        if x == "i":
            return "invalid"
        if x == "u":
            return "unsure"
        if x == "":
            return "valid"
        return x
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return "valid"
    return x

df["复核结果"] = df["复核结果"].apply(normalize)
print("规范化后分布:")
print(df["复核结果"].value_counts())

wb = load_workbook(SRC)
ws = wb["复核队列"]
header = [cell.value for cell in ws[1]]
col_idx = header.index("复核结果") + 1

VALID_FILL = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
INVALID_FILL = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
UNSURE_FILL = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")

for r in range(2, ws.max_row + 1):
    val = df.iloc[r - 2]["复核结果"]
    cell = ws.cell(row=r, column=col_idx)
    cell.value = val
    if val == "valid":
        cell.fill = VALID_FILL
    elif val == "invalid":
        cell.fill = INVALID_FILL
    elif val == "unsure":
        cell.fill = UNSURE_FILL
    cell.font = Font(bold=True)

wb.save(SRC)
print("已写入 Excel")

# 同步 CSV
csv_df = pd.read_csv(CSV)
for _, row in df.iterrows():
    rid = row["编号"]
    dec = row["复核结果"]
    csv_df.loc[csv_df["row_id"] == rid, "final_decision"] = dec
csv_df.to_csv(CSV, index=False, encoding="utf-8-sig")
print("已同步 CSV")
