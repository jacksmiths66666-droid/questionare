import pandas as pd
import openpyxl

wb = openpyxl.load_workbook('outputs/screening_v3/review_batch_A_431.xlsx')
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'Sheet: {sn}, rows={ws.max_row}, cols={ws.max_column}')
    # read header row
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
    print(f'Headers: {headers}')
    if sn == wb.sheetnames[0]:
        # Check last few rows for user input
        for r in range(max(2, ws.max_row-5), ws.max_row+1):
            row_id = ws.cell(row=r, column=2).value
            last_val = ws.cell(row=r, column=ws.max_column).value
            second_last = ws.cell(row=r, column=ws.max_column-1).value
            print(f'  Row{r}: id={row_id}, col14={second_last}, col15={last_val}')
